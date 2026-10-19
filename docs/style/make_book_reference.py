#!/usr/bin/env python3
"""Derive the full-book style docs/style/reference_book.docx from reference.docx by adding a
centered page number in the footer (hidden on the title page). Used by docs/build_book.py."""
from pathlib import Path
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
doc = docx.Document(str(HERE / "reference.docx"))
for sec in doc.sections:
    sec.different_first_page_header_footer = True
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)
    run.font.size = docx.shared.Pt(9)
out = HERE / "reference_book.docx"
doc.save(str(out))
print("wrote", out)
